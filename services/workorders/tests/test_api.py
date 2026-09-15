import io

import pytest
from openpyxl import Workbook, load_workbook

from cmms_common.contracts import assert_matches_contract
from workorders_api.models import WorkOrder, WorkOrderComment, WorkOrderPart

pytestmark = pytest.mark.django_db


def test_create_work_order_sets_requester_and_cached_asset(api_client, asset, admin_user, event_bus):
    response = api_client.post('/api/workorders/', {
        'equipment': asset.asset_id,
        'wo_type': 'CM',
        'summary': 'New work order',
        'description': 'Test description',
        'priority': 'high',
    }, format='json')
    assert response.status_code == 201, response.data
    data = response.json()
    assert data['requested_by'] == admin_user.id
    assert data['requester_name'] == 'Admin User'
    assert data['equipment'] == asset.asset_id
    assert data['equipment_code'] == 'EQ-001'
    assert data['equipment_name'] == 'Pump A'
    assert data['wo_code'].startswith('WO-')
    assert_matches_contract(data, 'workorders', 'WorkOrderSerializer')


def test_create_requires_supervisor(technician_client, asset):
    response = technician_client.post('/api/workorders/', {
        'equipment': asset.asset_id, 'summary': 'x',
    }, format='json')
    assert response.status_code == 403


def test_list_detail_update_delete(api_client, work_order):
    response = api_client.get('/api/workorders/')
    assert response.status_code == 200
    results = response.json()['results']
    assert len(results) == 1
    assert_matches_contract(results[0], 'workorders', 'WorkOrderListSerializer')

    response = api_client.get(f'/api/workorders/{work_order.id}/')
    assert response.status_code == 200
    assert response.json()['wo_code'] == work_order.wo_code

    response = api_client.patch(
        f'/api/workorders/{work_order.id}/',
        {'summary': 'Updated', 'maintenance_plan': 5},
        format='json',
    )
    assert response.status_code == 200
    assert response.json()['summary'] == 'Updated'
    assert response.json()['maintenance_plan'] == 5

    response = api_client.delete(f'/api/workorders/{work_order.id}/')
    assert response.status_code == 204
    assert not WorkOrder.objects.filter(id=work_order.id).exists()


def test_filters(api_client, work_order):
    work_order.assignee_id = 7
    work_order.save()
    assert len(api_client.get('/api/workorders/?assignee=7').json()['results']) == 1
    assert len(api_client.get('/api/workorders/?assignee=8').json()['results']) == 0
    assert len(api_client.get(f'/api/workorders/?equipment={work_order.equipment}').json()['results']) == 1
    assert len(api_client.get('/api/workorders/?status=closed').json()['results']) == 0
    assert len(api_client.get('/api/workorders/?search=leaking').json()['results']) == 1


def test_lifecycle_actions(api_client, work_order, admin_user, event_bus):
    url = f'/api/workorders/{work_order.id}'
    response = api_client.post(f'{url}/assign/', {'assignee_id': 7, 'assignee_name': 'Tech One'}, format='json')
    assert response.status_code == 200, response.data
    data = response.json()
    assert data['status'] == 'assigned'
    assert data['assignee'] == 7
    assert data['assignee_name'] == 'Tech One'
    assert data['assigned_by'] == admin_user.id

    response = api_client.post(f'{url}/assign/', {'assignee_id': 8}, format='json')
    assert response.status_code == 400

    response = api_client.post(f'{url}/start/')
    assert response.status_code == 200
    assert response.json()['status'] == 'in_progress'

    response = api_client.post(f'{url}/complete/', {}, format='json')
    assert response.status_code == 400

    response = api_client.post(f'{url}/complete/', {
        'actions_taken': 'Replaced seal', 'labor_hours': '2.5', 'parts_cost': '30.00',
    }, format='json')
    assert response.status_code == 200, response.data
    data = response.json()
    assert data['status'] == 'completed'
    assert data['completed_by'] == admin_user.id
    assert data['total_cost'] == '30.00'

    response = api_client.post(f'{url}/close/')
    assert response.status_code == 200
    data = response.json()
    assert data['status'] == 'closed'
    assert data['closed_by'] == admin_user.id
    assert_matches_contract(data, 'workorders', 'WorkOrderSerializer')


def test_assign_falls_back_to_id_as_name(api_client, work_order, event_bus):
    response = api_client.post(f'/api/workorders/{work_order.id}/assign/', {'assignee_id': 42}, format='json')
    assert response.status_code == 200
    assert response.json()['assignee_name'] == '42'


def test_my_orders_and_overdue(technician_client, work_order, technician_user):
    work_order.assignee_id = technician_user.id
    work_order.planned_end = '2000-01-01T00:00:00Z'
    work_order.save()
    response = technician_client.get('/api/workorders/my_orders/')
    assert response.status_code == 200
    assert [wo['id'] for wo in response.json()] == [work_order.id]
    response = technician_client.get('/api/workorders/overdue/')
    assert response.status_code == 200
    assert [wo['id'] for wo in response.json()] == [work_order.id]
    assert_matches_contract(response.json()[0], 'workorders', 'WorkOrderListSerializer')


def test_comments_and_parts(api_client, work_order, admin_user):
    response = api_client.post('/api/workorders/comments/', {
        'work_order': work_order.id, 'comment': 'Looks bad', 'is_internal': False,
    }, format='json')
    assert response.status_code == 201, response.data
    data = response.json()
    assert data['author'] == admin_user.id
    assert data['author_name'] == 'Admin User'
    assert_matches_contract(data, 'workorders', 'WorkOrderCommentSerializer')
    assert WorkOrderComment.objects.get(id=data['id']).author_id == admin_user.id
    assert len(api_client.get(f'/api/workorders/comments/?work_order={work_order.id}').json()['results']) == 1

    response = api_client.post('/api/workorders/parts/', {
        'work_order': work_order.id, 'part_code': 'P-1', 'part_name': 'Seal',
        'quantity': '2', 'unit': 'pc', 'unit_cost': '5.50',
    }, format='json')
    assert response.status_code == 201, response.data
    data = response.json()
    assert data['total_cost'] == '11.00'
    assert_matches_contract(data, 'workorders', 'WorkOrderPartSerializer')
    assert WorkOrderPart.objects.filter(work_order=work_order).count() == 1


def test_assign_serializer_contract():
    from workorders_api.serializers import WorkOrderAssignSerializer

    serializer = WorkOrderAssignSerializer(data={'assignee_id': 3})
    assert serializer.is_valid()
    assert_matches_contract(serializer.validated_data, 'workorders', 'WorkOrderAssignSerializer')


def test_template_import_export(api_client, asset, event_bus):
    response = api_client.get('/api/workorders/download_template/')
    assert response.status_code == 200
    template = load_workbook(io.BytesIO(response.content))
    headers = [cell.value.split('\n')[0].strip('*') for cell in template.active[1]]
    assert headers[0] == '设备编码'

    wb = Workbook()
    ws = wb.active
    ws.append(['设备编码', '工单摘要', '工单类型', '优先级'])
    ws.append(['EQ-001', 'Imported order', 'PM', 'low'])
    ws.append(['MISSING', 'Bad row', 'CM', 'low'])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    buffer.name = 'import.xlsx'
    response = api_client.post('/api/workorders/import_excel/', {'file': buffer}, format='multipart')
    assert response.status_code == 200, response.data
    assert response.json()['success_count'] == 1
    assert response.json()['error_count'] == 1
    imported = WorkOrder.objects.get(summary='Imported order')
    assert imported.equipment == asset.asset_id
    assert imported.equipment_code == 'EQ-001'
    assert imported.wo_type == 'PM'

    response = api_client.get('/api/workorders/export_csv/')
    assert response.status_code == 200
    exported = load_workbook(io.BytesIO(response.content)).active
    assert exported.cell(row=2, column=2).value == 'EQ-001'
    assert exported.cell(row=2, column=3).value == 'Pump A'


def test_health_and_metrics(api_client):
    assert api_client.get('/health/').status_code == 200
    assert api_client.get('/metrics').status_code == 200
