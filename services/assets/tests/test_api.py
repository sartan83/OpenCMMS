import io
from datetime import date, timedelta

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from cmms_common.contracts import assert_matches_contract

from assetregistry.models import Asset

ASSET_PAYLOAD = {
    'code': 'EQ-100',
    'name': 'CNC Mill',
    'process': 'Machining',
    'factory': 'F1',
    'workshop': 'W1',
    'line': 'L1',
    'station': 'S1',
    'status': 'active',
    'criticality': 'important',
}


@pytest.mark.django_db
def test_create_asset_sets_created_by_from_remote_user(api_client, admin_user, event_bus):
    response = api_client.post('/api/assets/', ASSET_PAYLOAD, format='json')
    assert response.status_code == 201, response.data
    data = response.json()
    assert data['created_by'] == admin_user.id
    assert data['created_by_name'] == admin_user.full_name
    assert data['location_display'] == 'F1 / W1 / L1 / S1'
    assert data['is_overdue'] is False
    assert data['level'] == 0
    asset = Asset.objects.get(id=data['id'])
    assert asset.created_by_id == admin_user.id
    assert_matches_contract(data, 'assets', 'AssetSerializer')


@pytest.mark.django_db
def test_list_retrieve_update_delete(api_client, asset_factory, event_bus):
    asset = asset_factory(code='EQ-001', name='Press', factory='F1')
    response = api_client.get('/api/assets/')
    assert response.status_code == 200
    results = response.json()['results']
    assert len(results) == 1
    assert_matches_contract(results[0], 'assets', 'AssetListSerializer')

    response = api_client.get(f'/api/assets/{asset.id}/')
    assert response.status_code == 200
    assert_matches_contract(response.json(), 'assets', 'AssetSerializer')

    response = api_client.patch(f'/api/assets/{asset.id}/', {'name': 'Hydraulic Press'}, format='json')
    assert response.status_code == 200
    assert response.json()['name'] == 'Hydraulic Press'

    response = api_client.delete(f'/api/assets/{asset.id}/')
    assert response.status_code == 204
    assert not Asset.objects.filter(id=asset.id).exists()


@pytest.mark.django_db
def test_non_admin_is_read_only(operator_client, asset_factory):
    asset_factory()
    assert operator_client.get('/api/assets/').status_code == 200
    response = operator_client.post('/api/assets/', ASSET_PAYLOAD, format='json')
    assert response.status_code == 403


def test_unauthenticated_is_rejected():
    from rest_framework.test import APIClient

    assert APIClient().get('/api/assets/').status_code in (401, 403)


@pytest.mark.django_db
def test_parent_cannot_be_self(api_client, asset_factory, event_bus):
    asset = asset_factory()
    response = api_client.patch(f'/api/assets/{asset.id}/', {'parent': asset.id}, format='json')
    assert response.status_code == 400
    assert 'parent' in response.json()


@pytest.mark.django_db
def test_tree_and_children(api_client, asset_factory):
    root = asset_factory(code='ROOT')
    child = asset_factory(code='CHILD', parent=root)
    asset_factory(code='INACTIVE', parent=root, status='inactive')

    response = api_client.get('/api/assets/tree/')
    assert response.status_code == 200
    tree = response.json()
    assert [node['code'] for node in tree] == ['ROOT']
    assert [node['code'] for node in tree[0]['children']] == ['CHILD']
    assert_matches_contract(tree[0], 'assets', 'AssetTreeSerializer')

    response = api_client.get(f'/api/assets/{root.id}/children/')
    assert response.status_code == 200
    assert {node['code'] for node in response.json()} == {'CHILD', 'INACTIVE'}

    response = api_client.get(f'/api/assets/{child.id}/')
    assert response.json()['level'] == 1


@pytest.mark.django_db
def test_overdue(api_client, asset_factory):
    asset_factory(code='LATE', next_maintenance_date=date.today() - timedelta(days=1))
    asset_factory(code='OK', next_maintenance_date=date.today() + timedelta(days=1))
    response = api_client.get('/api/assets/overdue/')
    assert response.status_code == 200
    assert [item['code'] for item in response.json()] == ['LATE']
    assert response.json()[0]['is_overdue'] is True


@pytest.mark.django_db
def test_filters_and_search(api_client, asset_factory):
    asset_factory(code='A1', factory='F1', vendor='Siemens')
    asset_factory(code='B2', factory='F2', vendor='Fanuc')
    response = api_client.get('/api/assets/', {'factory': 'F2'})
    assert [item['code'] for item in response.json()['results']] == ['B2']
    response = api_client.get('/api/assets/', {'search': 'Siemens'})
    assert [item['code'] for item in response.json()['results']] == ['A1']


@pytest.mark.django_db
def test_download_template_and_export(api_client, asset_factory):
    asset_factory(code='EXP-1', name='Exported')
    response = api_client.get('/api/assets/download_template/')
    assert response.status_code == 200
    assert response['Content-Type'].startswith('application/vnd.openxmlformats')
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert wb.active.cell(row=2, column=1).value == 'EQ-001'

    response = api_client.get('/api/assets/export_csv/')
    assert response.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert wb.active.cell(row=2, column=1).value == 'EXP-1'


@pytest.mark.django_db
def test_import_excel_creates_assets(api_client, admin_user, event_bus, django_capture_on_commit_callbacks):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['设备编码*\n(必填，唯一)', '设备名称*\n(必填)', '工厂', '投用日期\n(格式: YYYY-MM-DD)'])
    ws.append(['IMP-1', 'Imported one', 'F1', '2024-01-15'])
    ws.append(['IMP-2', 'Imported two', 'F2', None])
    ws.append([None, 'Missing code', None, None])
    output = io.BytesIO()
    wb.save(output)
    upload = SimpleUploadedFile('assets.xlsx', output.getvalue())

    with django_capture_on_commit_callbacks(execute=True):
        response = api_client.post('/api/assets/import_excel/', {'file': upload}, format='multipart')
    assert response.status_code == 200, response.data
    body = response.json()
    assert body['success_count'] == 2
    assert body['error_count'] == 1
    imported = Asset.objects.get(code='IMP-1')
    assert imported.start_date == date(2024, 1, 15)
    assert imported.created_by_id == admin_user.id
    assert imported.created_by_name == admin_user.full_name
    assert [e.type for e in event_bus.published] == ['asset.created', 'asset.created']
    assert {e.payload['code'] for e in event_bus.published} == {'IMP-1', 'IMP-2'}


@pytest.mark.django_db
def test_import_csv(api_client, event_bus):
    csv_content = 'code,name,factory\nCSV-1,From CSV,F9\n'.encode('utf-8-sig')
    upload = SimpleUploadedFile('assets.csv', csv_content)
    response = api_client.post('/api/assets/import_excel/', {'file': upload}, format='multipart')
    assert response.status_code == 200, response.data
    assert Asset.objects.get(code='CSV-1').factory == 'F9'


@pytest.mark.django_db
def test_import_rejects_missing_or_bad_file(api_client):
    assert api_client.post('/api/assets/import_excel/', {}, format='multipart').status_code == 400
    upload = SimpleUploadedFile('assets.txt', b'nope')
    assert api_client.post('/api/assets/import_excel/', {'file': upload}, format='multipart').status_code == 400


@pytest.mark.django_db
def test_health_and_metrics():
    from rest_framework.test import APIClient

    client = APIClient()
    assert client.get('/health/').status_code == 200
    assert client.get('/metrics').status_code == 200
